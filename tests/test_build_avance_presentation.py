"""Presentation builder works with or without an avance template."""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import pytest

pptx = pytest.importorskip('pptx')
from pptx import Presentation

SCRIPT = Path(__file__).resolve().parents[1] / 'scripts' / 'build_avance_presentation.py'


def _load_script():
    spec = importlib.util.spec_from_file_location('build_avance_presentation', SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def _unsupported_items(path: Path) -> Path:
    items = [
        {
            'construct': 'egocentric_encoding',
            'status': 'unsupported',
            'question_style': 'concise',
            'distractor_rationale': {'reason': 'test_blocker'},
        }
    ]
    dest = path / 'house_test' / 'items_house_test.json'
    dest.parent.mkdir(parents=True)
    dest.write_text(json.dumps({'n_items': len(items), 'items': items}))
    return dest.parent.parent


@pytest.fixture(scope='module')
def slides():
    return _load_script()


def test_blank_layout_on_default_deck(slides):
    prs = Presentation()
    layout = slides.blank_layout(prs)
    added = prs.slides.add_slide(layout)
    assert added is not None
    assert str(getattr(layout, 'name', '')).lower() == 'blank'


def test_build_without_template(tmp_path, slides):
    items_root = _unsupported_items(tmp_path)
    out = tmp_path / 'examples.pptx'
    slides.main([
        '--items-json', str(items_root),
        '--output', str(out),
    ])
    assert out.is_file()
    prs = Presentation(str(out))
    assert len(prs.slides) >= 4


def test_build_with_template(tmp_path, slides):
    items_root = _unsupported_items(tmp_path)
    template = tmp_path / 'template.pptx'
    seed = Presentation()
    layout = slides.blank_layout(seed)
    for _ in range(8):
        seed.slides.add_slide(layout)
    seed.save(str(template))

    out = tmp_path / 'from_template.pptx'
    slides.main([
        '--template', str(template),
        '--items-json', str(items_root),
        '--output', str(out),
    ])
    prs = Presentation(str(out))
    assert len(prs.slides) > 8
